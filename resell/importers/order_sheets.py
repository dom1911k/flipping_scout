"""Best-effort parser for the messy "Order #N" tabs.

Each tab holds one or more blocks. A block starts with a header row (has a "Date" cell plus
amount columns); rows below it are line items until the layout changes. Tabs with no header
at all (Order #6/#7) use DEFAULT_LAYOUT: link | item | date | qty | RMB | EUR | fee | ship | total.
Totals rows (numbers, no item) are skipped silently; anything else that doesn't fit is
written to parse_issues rather than guessed.
"""
from __future__ import annotations

import re
from datetime import date, datetime

import openpyxl
import pandas as pd

from .sales_sheet import _num, _text

# header text (lower-case) -> (field, currency hint)
HEADER_MAP: dict[str, tuple[str, str | None]] = {
    "article": ("item", None),
    "date": ("date", None),
    "amount purchased": ("qty", None),
    "amount": ("qty", None),
    "amount (yen)": ("unit_local", "JPY"),
    "amount (usd)": ("unit_local", "USD"),
    "amount yuan": ("unit_local", "CNY"),
    "add. fee": ("fee_local", None),
    "add costs": ("fee_eur", None),
    "amount (eur at the time)": ("eur_goods", None),
    "overall euro": ("eur_goods", None),
    "vat final price": ("vat_incl_eur", None),
    "shipping per item": ("shipping_eur", None),
    "shipping": ("shipping_eur", None),
    "add delivery fee": ("shipping2_eur", None),
    "final price": ("total_eur", None),
    "final price incl shipping": ("total_eur", None),
    "overall": ("total_eur", None),
}
IGNORED_HEADERS = {"money spent on repair tools", "overall spent", "overall earned"}

DEFAULT_LAYOUT = {1: "link", 2: "item", 3: "date", 4: "qty", 5: "unit_local", 6: "eur_goods",
                  7: "fee_eur", 8: "shipping_eur", 9: "total_eur", 10: "note", 11: "note"}
DEFAULT_CURRENCY = "CNY"
NOTE_SPAN = 3  # text in up to N unlabelled columns right of the block counts as a line note


def _is_date(v) -> bool:
    return isinstance(v, (datetime, date))


def _parse_header(cells: dict[int, object]) -> tuple[dict[int, str], str] | None:
    labels = {c: str(v).strip().lower() for c, v in cells.items() if isinstance(v, str)}
    if "date" not in labels.values():
        return None
    layout: dict[int, str] = {}
    currency = None
    for col, text in labels.items():
        if text in IGNORED_HEADERS:
            continue
        if text in HEADER_MAP:
            field, cur = HEADER_MAP[text]
            layout[col] = field
            currency = currency or cur
    if not {"unit_local", "eur_goods", "total_eur"} & set(layout.values()):
        return None
    # USD parts table: "Overall" is the USD line total, the EUR total is "Final price incl shipping"
    if currency == "USD":
        for col, text in labels.items():
            if text == "overall":
                layout[col] = "local_total"
    if "eur_goods" in layout.values() and currency is None:
        currency = "CNY" if "unit_local" in layout.values() else "EUR"
    date_col = next(c for c, f in layout.items() if f == "date")
    if "item" not in layout.values() and date_col > 1:
        layout[date_col - 1] = "item"  # Order #5 has blocks with no "Article" label
    item_col = next(c for c, f in layout.items() if f == "item")
    if item_col > 1 and (item_col - 1) not in layout:
        layout[item_col - 1] = "link"  # seller link / category label left of the item
    return layout, currency or "CNY"


def _derive(rec: dict, currency: str) -> dict:
    qty = rec.get("qty") or 1.0
    unit = rec.get("unit_local")
    eur = rec.get("eur_goods")
    out = {
        "qty": qty,
        "currency": currency,
        "unit_price_local": unit,
        "eur_goods": eur,
        "fee_eur": rec.get("fee_eur"),
        "shipping_eur": (rec.get("shipping_eur") or 0) + (rec.get("shipping2_eur") or 0) or rec.get("shipping_eur"),
        "vat_incl_eur": rec.get("vat_incl_eur"),
        "total_eur": rec.get("total_eur"),
        "eff_rate": None,
    }
    if unit and eur:
        local_total = unit * qty + (rec.get("fee_local") or 0)
        out["eff_rate"] = local_total / eur
    return out


def read_orders(wb: openpyxl.Workbook, prefix: str) -> tuple[pd.DataFrame, list[dict]]:
    rows: list[dict] = []
    issues: list[dict] = []
    for ws in wb.worksheets:
        if not ws.title.startswith(prefix):
            continue
        layout, currency = dict(DEFAULT_LAYOUT), DEFAULT_CURRENCY
        block = 0
        side_cells: list[str] = []
        for excel_row in ws.iter_rows():
            cells = {c.column: c.value for c in excel_row if c.value is not None and str(c.value).strip() != ""}
            if not cells:
                continue
            r = excel_row[0].row
            header = _parse_header(cells)
            if header:
                layout, currency = header
                block += 1
                continue
            by_field: dict[str, object] = {}
            notes, extra_nums, side = [], [], []
            last_col = max(layout)
            for col, v in cells.items():
                field = layout.get(col)
                if field == "note" or (field is None and last_col < col <= last_col + NOTE_SPAN and isinstance(v, str)):
                    notes.append(_text(v))
                elif field is None and col <= last_col + NOTE_SPAN:
                    extra_nums.append(f"{openpyxl.utils.get_column_letter(col)}{r}={v}")
                elif field is None:
                    side.append(f"{openpyxl.utils.get_column_letter(col)}{r}={_text(v)}")
                else:
                    by_field[field] = v
            side_cells += side
            item = _text(by_field.get("item"))
            if item and not re.search(r"[^\W\d_]", item):
                item = None  # a bare number in the article column is not a name
            date_v = by_field.get("date")
            numeric = {k: _num(by_field.get(k)) for k in
                       ("qty", "unit_local", "fee_local", "fee_eur", "eur_goods", "vat_incl_eur",
                        "shipping_eur", "shipping2_eur", "total_eur", "local_total")}
            has_amount = any(numeric[k] is not None for k in ("unit_local", "eur_goods", "total_eur"))
            if not item:
                if has_amount and not _is_date(date_v):
                    continue  # totals row
                if has_amount or _is_date(date_v):
                    issues.append({"sheet": ws.title, "cells": f"row {r}", "raw": repr(cells),
                                   "reason": "amounts but no item name"})
                continue
            if not has_amount:
                issues.append({"sheet": ws.title, "cells": f"row {r}", "raw": repr(cells),
                               "reason": "item without any amount (label or sub-heading?)"})
                continue
            rec = _derive(numeric, currency)
            if not _is_date(date_v):
                issues.append({"sheet": ws.title, "cells": f"row {r}", "raw": item,
                               "reason": "line item has no date - imported with empty date"})
            if extra_nums:
                issues.append({"sheet": ws.title, "cells": ", ".join(extra_nums), "raw": item,
                               "reason": "unlabelled numbers next to a line item - ignored"})
            link = _text(by_field.get("link"))
            if link and "://" not in link and len(link) <= 10:
                notes.insert(0, link)  # short label such as "GBA" left of the item
                link = None
            rows.append({
                "sheet": ws.title,
                "sheet_row": r,
                "block": block,
                "item": item,
                "buy_date": date_v.date() if isinstance(date_v, datetime) else (date_v if _is_date(date_v) else None),
                "link": link,
                "note": "; ".join(n for n in notes if n) or None,
                **rec,
            })
        if side_cells:
            issues.append({"sheet": ws.title, "cells": f"{len(side_cells)} cells",
                           "raw": " | ".join(side_cells[:12]) + (" ..." if len(side_cells) > 12 else ""),
                           "reason": "side table outside the order columns - not imported"})
    df = pd.DataFrame(rows)
    df["excluded"] = None
    _flag_copied_blocks(df, issues)
    return df, issues


def _flag_copied_blocks(df: pd.DataFrame, issues: list[dict]) -> None:
    """Some tabs repeat a block with other prices (what-if copies). If most (item, date) pairs of a
    block already appeared in an earlier block of the same tab, exclude it from statistics."""
    for sheet, g in df.groupby("sheet"):
        seen: set = set()
        for blk, b in g.groupby("block", sort=True):
            keys = set(zip(b["item"].str.lower(), b["buy_date"]))
            if len(keys) >= 3 and len(keys & seen) / len(keys) >= 0.6:
                df.loc[b.index, "excluded"] = "copy of earlier block"
                issues.append({"sheet": sheet, "cells": f"rows {b.sheet_row.min()}-{b.sheet_row.max()}",
                               "raw": f"{len(b)} lines", "reason": "block repeats an earlier block with "
                               "different amounts (what-if copy?) - imported but excluded from stats"})
            seen |= keys
